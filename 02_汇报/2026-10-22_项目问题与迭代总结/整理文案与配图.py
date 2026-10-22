from pathlib import Path
import shutil, zipfile, json
from PIL import Image, ImageOps, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

OUT=Path(__file__).parent
ROOT=OUT.parent.parent
REPO=Path('E:/课程作业/milk-github-sync')
ATT=Path('E:/作品集/抖音ai创作大赛26/.codex-remote-attachments/01a09d93-5828-7951-8f50-b4e4bdacf50a')
ASSETS=OUT/'配图素材'; ASSETS.mkdir(exist_ok=True)

core=[
('发射器与接收器的连接问题',
 '灯亮、配对成功，不代表机器人正在执行我们的程序。',
 '前期反复出现板子有灯、遥控却没反应；重新配对后还发生过接收板文件被清空。拳头版又遇到原厂“叉车”程序仍在运行，造成单关节乱动和电机自转。',
 '保留官方无线通信，分开检查配对、摇杆数据和动作输出；禁用接收端原厂自动控制入口，安装自己的上电启动程序，上传后校验文件并重启。',
 '两台均已部署持续遥控入口，并有现场恢复控制的反馈。灯色只作为提示，不能单独判断是否配对成功或程序正确。'),
('机器人结构牢固性问题',
 '能拼起来，还要保证连接可靠、活动关节不卡死。',
 '攻击模块与身体需要新增连接；拳头位置调整后，原安装孔也随之错位。曲柄、连杆和滑架既要连接可靠，又必须保留相对转动。',
 '将连接圈、连接梁与开放底座整合，匹配现有身体；直线滑架使用两根金属导杆，活动铰点采用螺钉、垫片与防松螺母，避免把活动关节拧成死连接。',
 '完成了装配设计及机构运动验证；抗冲击、防松和长期对战耐久性还需要继续测试。'),
('机器人的重心问题',
 '攻击模块越往前，出拳越明显，但整机越容易失衡。',
 '最初预览中拳头偏斜、机构整体过于靠前，拳头在部分行程还被侧面结构遮挡。外形、出拳距离与平衡不能分开调整。',
 '先在整机预览中调整拳头朝向、位置与腕杆长度，将攻击机构相对身体后移22毫米，再重做与身体匹配的连接底座，保留已打印的下半身。',
 '已完成位置与接口调整；后移是针对失衡风险的设计改进，尚未测出整机真实重心，也未证明所有攻击动作下都能保持稳定。'),
('机器人行走姿态控制问题',
 '四个舵机分别会动，不等于机器人会走、会转、会倒退。',
 '先后出现原地踏步、左右杆转向相似，以及同一根杆前推和后拉都沿前进弧线走。机械中位、镜像安装、髋踝相位和方向逻辑都会影响结果。',
 '先将四只位置舵机归到90度再安装，逐路确认方向；参考Otto连续正弦步态协调髋踝。再把前后方向与左右步幅分配拆开，换向前先收步，并让两台共用同一套行走逻辑。',
 '基础行走已有落地实测；最新统一版已于10月21日上传两台。单杆前后弧线修复有代码回归验证，但最新转向轨迹仍需现场对照确认。')]

extra=[
('测试结束后遥控失效', '把临时测试改成真正的持续遥控。',
 '最初控制窗口只有几十秒或一两分钟，到时程序主动退出，看起来像突然断连。后来增加上电自启和持续运行入口，保留独立的限时诊断工具。',
 '持续遥控已获用户现场确认；正常使用不再受测试倒计时影响。'),
('接线松动造成无响应', '先看原始输入，再决定要不要改代码。',
 '10月21日，右杆松手和后拉到底时ch0都约为3750，程序始终等不到中位；重插摇杆线后，用户确认完全恢复控制。随后巴掌攻击杆读数正常、直驱舵机也不转，检查连接并重新供电后用户反馈恢复转动。',
 '右杆故障已确认是接线松动；后一次巴掌故障未单独隔离到具体焊点或供电部位，不写成已确诊某一元件损坏。'),
('舵机与电机的选择', '腿部要能定位，巴掌要能连续旋转，拳头要能持续往复。',
 '早期腿部混入连续旋转舵机，无法像位置舵机一样保持指定角度；四个腿部舵机口用满后，顶部攻击也需要新的输出路径。最终腿部使用四只180度位置舵机，巴掌连续舵机接D1，拳头电机接M1并通过减速和曲柄滑块实现往复。',
 '两种攻击均已有实机动作确认；不能把连续舵机的1500µs停止信号理解成90度位置。'),
('攻击手感与控制语义', '玩家推住就持续攻击，松手就停。',
 '拳头曾持续推杆几秒后自动停止，原因是单次攻击时长限制；取消该限制后实现持续出拳，并加入上推慢、下推快。巴掌早期1400/1600µs带不动，扩大偏差后能转；最新按用户选择设为900/2100µs，回中1500µs停止。',
 '持续出拳及两档速度获现场确认。巴掌最高指令短测曾确认更快、机构正常；最新900/2100µs已上传，不将指令档位当作实际功率测量。'),
('套件更换与接口适配', '从套件默认能力出发，重新规划机械和程序。',
 '项目从CyberBrick 2切换到CyberBrick 1，原先对舵机外形、数量和接口的假设需要重查。我们核对了180度与连续旋转舵机、GPIO和连接形式，保留Otto短腿方案，按实际组件调整顶部模块与控制入口。',
 'CB1双机方案已进入实物运行；规格核对、接口标注和统一安装顺序成为复刻第二台的重要依据。'),
('玩法与场景的适配', '机器人能完成的动作，决定场地和目标怎么设计。',
 '双足机器人转向、通行与击打高度受限，场地不能只按效果图设计。地图改为单层平地并调整障碍与通行空间；球靶方案按两机模型的攻击高度重新设计，优先打印，并用可调高度和磁吸脱落形成可见反馈。',
 '用户已确认竞技场搭建完成；球靶已交付打印与装配文件，实际脱落手感、磁力和耐打性不写成已验收。')]

assets=[
('01_接收扩展板正面_现场照片.jpg',ATT/'68e87115-cae1-4de8-bdfc-6e447bc61fe0/2-Photo-2.jpg','实物照片','接口扩展：四路舵机口与电机口。不能用它证明无线配对状态。'),
('02_板卡背面_现场照片.jpg',ATT/'68e87115-cae1-4de8-bdfc-6e447bc61fe0/1-Photo-1.jpg','实物照片','板卡与硬件排查过程，配合接口说明使用。'),
('03_舵机接口不匹配_现场照片.jpg',ATT/'763f91a8-5f86-4196-a0e4-7bba41f46830/1-Photo-1.jpg','实物照片','标准舵机头与小型白色插头的差异。'),
('04_电机与舵机对比_现场照片.jpg',ATT/'f4761830-4aff-4445-9262-d7a1ac32300c/1-Photo-1.jpg','实物照片','说明两种执行器的选择，不表示裸电机能直接完成出拳。'),
('05_减速齿轮包_现场照片.jpg',ATT/'c155bebf-af22-414e-ad86-e7ccae083f0f/1-Photo-1.jpg','实物照片','拳头机构需要减速组件及相应安装结构。'),
('06_拳头俯视检查_模型截图.png',ROOT/'04_结构与硬件/拳头_整体比例预览_v02/check_top.png','模型截图','展示朝向、位置与整机比例检查，不能当实物照片。'),
('07_拳头位置调整_模型截图.png',ROOT/'04_结构与硬件/拳头_整体比例预览_v03/check_adjusted.png','模型截图','展示整机预览中的攻击模块位置调整。'),
('08_拳头模块装配标注_模型图.png',ROOT/'04_结构与硬件/拳头_全套打印包_2026-09-29/成品装配图_零件标注.png','装配模型图','对应底座、曲柄、连杆、滑架、拳头等连接关系。'),
('09_巴掌连接剖视_模型图.png',ROOT/'04_结构与硬件/Otto_75mm巴掌腕部接口_v02/腕部连接_装配拆解剖视.png','装配模型图','说明巴掌与横臂、腕部接口的连接设计。'),
('10_球靶调高与锁紧_模型图.png',ROOT/'06_游戏设计/打印脱落球靶_v09/制造初样_v01/04_总装与检查/01_总装与调高锁紧.png','装配模型图','说明按打击高度设计目标和锁紧结构；不代表实打验收。')]
catalog=['# 配图说明','', '照片保留原貌；模型截图与装配图只说明设计过程，不替代实物验证。','']
valid=[]
for name,src,kind,use in assets:
 if src.exists():
  shutil.copy2(src,ASSETS/name)
  valid.append((name,src,kind,use))
  catalog += [f'## {name}',f'- 类型：{kind}',f'- 适合说明：{use}',f'- 原始文件：{src}','']
(OUT/'配图说明.md').write_text('\n'.join(catalog),encoding='utf-8')

fontpath='C:/Windows/Fonts/msyh.ttc'
font=ImageFont.truetype(fontpath,21); small=ImageFont.truetype(fontpath,17)
sheet=Image.new('RGB',(1400,5*330+95),'#f4f4f4');draw=ImageDraw.Draw(sheet)
draw.text((30,20),'milk 项目调试配图索引',font=ImageFont.truetype(fontpath,29),fill='#111')
for i,(name,src,kind,use) in enumerate(valid):
 x=20+(i%2)*690;y=90+(i//2)*330
 draw.rounded_rectangle((x,y,x+670,y+310),12,fill='white')
 im=ImageOps.exif_transpose(Image.open(src)).convert('RGB')
 im.thumbnail((640,245))
 sheet.paste(im,(x+(670-im.width)//2,y+8+(245-im.height)//2))
 draw.text((x+12,y+260),name.rsplit('.',1)[0],font=small,fill='#222')
 draw.text((x+12,y+285),kind,font=small,fill='#555')
sheet.save(OUT/'配图总览.jpg',quality=91)

intro='我们围绕无线控制、机械装配和行走协同反复迭代，将两个能动的原型逐步调整为可以持续遥控的桌面对战机器人。最关键的经验是把“输入是否正确”“程序是否持续运行”“机构是否真正完成动作”分开验证。'
onepage='''发射器与接收器连接：灯亮不等于程序运行，统一启动入口并逐段验证无线输入与动作输出。
结构牢固性：重做底座与连接关系，用金属导杆和活动铰点实现可靠装配。
整机重心：攻击模块后移，重新调整拳头前伸与身体连接，兼顾打击距离和平衡。
行走姿态：从单关节动作转向四关节协同，修正髋踝相位、单杆前后方向和两机参数。'''
speech='刚开始，我们以为难点是让舵机动起来，真正装好以后才发现，输入、程序和结构必须一起配合。板子亮灯不代表自己的程序在持续运行；四只舵机都动，也可能只是在原地踏步。我们一方面拆开排查配对、启动入口和接线，另一方面参考Otto的连续步态重做四关节协同。加上攻击模块后，又通过整机预览调整拳头位置、后移机构并重做连接底座。现在两台已经具备持续遥控与各自攻击功能，后续还要重点验证转向一致性、对战抗冲击和长时间运行。'
md=['# milk组 项目问题与迭代总结','', '2026年10月22日｜PPT文案与配图参考','',intro,'','## 一页精简版','',onepage,'','## 四个核心问题 大字与小字文案','']
for title,big,problem,change,result in core:
 md += [f'### {title}',f'**大字：{big}**','',f'**小字·问题：** {problem}',f'**小字·调整：** {change}',f'**小字·结果：** {result}','']
md += ['## 可选补充问题','']
for title,big,detail,result in extra:
 md += [f'### {title}',f'**大字：{big}**','',f'**小字：** {detail}',f'**当前结果：** {result}','']
md += ['## 建议放进PPT的方式','',
 '建议用两页讲问题：第一页放连接、牢固性、重心、步态四个核心问题；第二页放测试超时、接线松动和攻击手感三个典型小案例。每项保留一句大字和两三行小字，其余内容作为讲稿。',
 '如果只有一页空间，直接用上方“一页精简版”，配拳头装配图和接口不匹配照片。结构牢固性与重心用“已调整、继续验证”，不写“彻底解决”。','',
 '## 约一分钟讲稿','',speech,'','## 配图与证据说明','',
 '配图素材共10张，包括5张历史实物照片与5张模型截图/装配图。另有配图总览和逐图说明。未找到可直接证明本机配对成功或前后转弯修复的历史界面截图，建议这两项使用已有文字事实，或后补现场操作视频。',
 '10月20日的“PC通过、待上传”属于历史状态；10月21日两台均已上传，巴掌接线恢复后用户确认完全恢复控制。上传和测试通过不等于所有转向轨迹、冲击强度和长期运行均已验收。','',
 '## 记录依据','',
 '- 2026-09-24遥控无动作公开资料复核与执行交接：区分通信、启动、执行和测试超时。',
 '- 2026-09-24深夜原地踏步卡点交接及后续9月26日验收记录：步态与转向问题的演变。',
 '- 2026-10-13拳头M1持续转动修复记录及10月17日基本完成交接：原厂程序、持续攻击和持续遥控。',
 '- 2026-10-20拳头单杆倒退修复报告及双机统一控制交接：方向与步幅分配、统一控制。',
 '- 2026-10-21两台上传记录与本次现场反馈：接线松动、回中使能、巴掌速度档位。',
 '- 拳头整体比例预览v03及全套打印包装配文件：后移22毫米、重做底座和活动连接。','']
text='\n'.join(md)
(OUT/'项目问题与迭代总结_PPT文案.md').write_text(text,encoding='utf-8')
(OUT/'项目问题与迭代总结_可直接复制.txt').write_text(text,encoding='utf-8-sig')

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=doc.styles[name];s.font.name='Microsoft YaHei';s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
 s.font.size=Pt(10 if name=='Normal' else 22 if name=='Title' else 15 if name=='Heading 1' else 12)
doc.styles['Normal'].paragraph_format.space_after=Pt(3)
doc.styles['Normal'].paragraph_format.line_spacing=1.05
for name in ['Heading 1','Heading 2']:
 doc.styles[name].paragraph_format.space_before=Pt(7)
 doc.styles[name].paragraph_format.space_after=Pt(3)
def para(t,size=10,bold=False):
 p=doc.add_paragraph();r=p.add_run(t);r.bold=bold;r.font.size=Pt(size);return p
doc.add_paragraph('milk组项目问题与迭代总结','Title')
para('PPT大字与小字文案  2026年10月22日',10)
para(intro)
doc.add_heading('四个核心问题',1)
for title,big,problem,change,result in core:
 doc.add_heading(title,2);para(big,11,True)
 para('问题  '+problem,9);para('调整  '+change,9);para('结果  '+result,9)
doc.add_heading('补充案例',1).paragraph_format.page_break_before=True
for title,big,detail,result in extra:
 doc.add_heading(title,2);para(big,10,True);para(detail,9);para('结果  '+result,9)
doc.add_heading('精简版与汇报讲稿',1).paragraph_format.page_break_before=True
para('一页PPT可直接使用',11,True)
for line in onepage.splitlines():para(line)
doc.add_heading('约一分钟讲稿',2);para(speech)
doc.add_heading('讲述边界',2)
para('已确认：基础行走、持续遥控和两种攻击均有现场反馈；两机统一代码已上传。仍待验证：最新转弯轨迹、实际重心、抗冲击耐久与长时间供电。',10)
para('不要把模型图当成实物照片，也不要把“板端日志完成”写成“机器人实测完成”。',10)
doc.add_heading('配图建议',2)
para('接口问题用01、03号实物照片；结构与重心用06、07、08号模型截图；攻击机构选型用04、05号照片。球靶设计用10号模型图。')
para('素材目录包含10张原图、配图总览及逐图说明。未找到可直接证明本机配对或转弯修复的历史界面截图。',9)
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
for p in doc.paragraphs:
 for border in list(p._p.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.save(OUT/'项目问题与迭代总结_PPT文案.docx')
with zipfile.ZipFile(OUT/'PPT问题总结与配图_发组员.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in [OUT/'项目问题与迭代总结_PPT文案.md',OUT/'项目问题与迭代总结_可直接复制.txt',OUT/'项目问题与迭代总结_PPT文案.docx',OUT/'配图说明.md',OUT/'配图总览.jpg',*ASSETS.iterdir()]:
  z.write(p,p.relative_to(OUT))
print(json.dumps({'images':len(valid),'output':str(OUT)},ensure_ascii=False))
